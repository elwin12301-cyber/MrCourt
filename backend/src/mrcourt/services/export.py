"""Draft export to Word and PDF. Runs only when a user asks for it."""

import io
import re
from pathlib import Path

from docx import Document
from fpdf import FPDF

from mrcourt.models import Draft

REVIEW_NOTE = "AI-generated draft. Review before use."

_FONT_CANDIDATES = [
    r"C:\Windows\Fonts\arial.ttf",
    r"C:\Windows\Fonts\segoeui.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
]
_ASCII_FALLBACK = {
    "₹": "Rs. ", "“": '"', "”": '"', "‘": "'", "’": "'", "—": "-", "–": "-", "·": "-", "…": "...",
}


def _unicode_font() -> str | None:
    return next((p for p in _FONT_CANDIDATES if Path(p).exists()), None)


def _latin1(text: str) -> str:
    for char, replacement in _ASCII_FALLBACK.items():
        text = text.replace(char, replacement)
    return text.encode("latin-1", "replace").decode("latin-1")


def safe_filename(title: str, extension: str) -> str:
    stem = re.sub(r"[^\w\- ]+", "", title).strip().replace(" ", "_") or "draft"
    return f"{stem}.{extension}"


def to_docx(draft: Draft) -> bytes:
    document = Document()
    document.add_heading(draft.title, level=1)
    for block in draft.body.split("\n\n"):
        document.add_paragraph(block)
    note = document.add_paragraph()
    note.add_run(REVIEW_NOTE).italic = True
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def to_pdf(draft: Draft) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    font_path = _unicode_font()
    if font_path:
        pdf.add_font("Body", "", font_path)
        family, prepare = "Body", (lambda s: s)
    else:  # no Unicode font on this machine: fall back to the built-in font, transliterating
        family, prepare = "Helvetica", _latin1

    pdf.set_font(family, size=16)
    pdf.multi_cell(0, 9, prepare(draft.title), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font(family, size=12)
    for block in draft.body.split("\n\n"):
        pdf.multi_cell(0, 7, prepare(block), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
    pdf.set_font(family, size=9)
    pdf.multi_cell(0, 6, prepare(REVIEW_NOTE), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())
